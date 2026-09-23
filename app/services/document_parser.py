"""Parsing des documents importés par l'analyste.

Formats : PDF, TXT, DOCX, Markdown, JSON, CSV.

Règles prompt injection : le contenu extrait est toujours des DONNÉES,
jamais des instructions exécutées par les agents. Le contenu est enveloppé
dans des marqueurs `<<<DEBUT_DOCUMENT_NON_FIABLE>>>` / `<<<FIN_DOCUMENT_NON_FIABLE>>>`
lors de la construction des prompts.
"""

from pathlib import Path

from app.services import config

# Extensions autorisées (configurables via .env : ALLOWED_EXTENSIONS)
ALLOWED_EXTENSIONS: set[str] = set(config.ALLOWED_EXTENSIONS)

# Taille maximale par défaut (MB) — configurable via MAX_FILE_SIZE_MB
DEFAULT_MAX_FILE_SIZE_MB = config.MAX_FILE_SIZE_MB

# Marqueurs délimitant le contenu non fiable des documents.
DOCUMENT_UNTRUSTED_BEGIN = "<<<DEBUT_DOCUMENT_NON_FIABLE>>>"
DOCUMENT_UNTRUSTED_END = "<<<FIN_DOCUMENT_NON_FIABLE>>>"


class DocumentParseError(ValueError):
    """Document refusé (type non autorisé, taille excessive, encodage…)."""


def validate_upload(filename: str, size_bytes: int, max_mb: int = DEFAULT_MAX_FILE_SIZE_MB) -> str:
    """Valide un fichier importé avant traitement.

    Retourne l'extension en minuscules si le fichier est accepté,
    lève DocumentParseError sinon.
    """
    if not filename or "." not in filename:
        raise DocumentParseError("nom de fichier invalide")
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentParseError(
            f"extension '{extension}' non autorisée "
            f"(autorisées : {', '.join(sorted(ALLOWED_EXTENSIONS))})"
        )
    if size_bytes > max_mb * 1024 * 1024:
        raise DocumentParseError(f"fichier trop volumineux (max {max_mb} MB)")
    return extension


def wrap_untrusted(content: str) -> str:
    """Enferme le contenu documentaire dans des marqueurs non fiables."""
    return f"{DOCUMENT_UNTRUSTED_BEGIN}\n{content}\n{DOCUMENT_UNTRUSTED_END}"


# ----------------------------------------------------------------------
# Extraction texte brut par extension
# ----------------------------------------------------------------------


def _extract_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("encodage non UTF-8 pris en charge") from exc


def _extract_pdf(content: bytes) -> str:
    """Extraction du texte d'un PDF (pas d'OCR — note cette limite)."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("pypdf non installé") from exc
    try:
        reader = PdfReader(__import__("io").BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # PDF malformé ou protégé
        raise DocumentParseError(f"PDF illisible : {exc}") from exc
    return "\n".join(page for page in pages if page.strip())


def _extract_docx(content: bytes) -> str:
    """Extraction du texte d'un DOCX (paragraphes et tableaux)."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("python-docx non installé") from exc
    try:
        document = Document(__import__("io").BytesIO(content))
    except Exception as exc:
        raise DocumentParseError(f"DOCX illisible : {exc}") from exc
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return "\n".join(parts)


_PARSERS = {
    "pdf": _extract_pdf,
    "txt": _extract_txt,
    "md": _extract_txt,
    "json": _extract_txt,  # le JSON est transmis tel quel au LLM
    "csv": _extract_txt,
    "docx": _extract_docx,
}


def parse(filename: str, content: bytes, max_mb: int = DEFAULT_MAX_FILE_SIZE_MB) -> str:
    """Parse un document en texte brut selon son extension."""
    extension = validate_upload(filename, len(content), max_mb)
    try:
        return _PARSERS[extension](content)
    except DocumentParseError:
        raise
    except Exception as exc:  # défense en profondeur
        raise DocumentParseError(f"échec d'extraction de '{filename}' : {exc}") from exc


def load_documents(
    files: list[tuple[str, bytes]],
    max_mb: int = DEFAULT_MAX_FILE_SIZE_MB,
) -> list[dict]:
    """Prépare la liste de documents analysables.

    Entrée : liste de (nom_fichier, contenu_binaire).
    Sortie : [{"id": "DOC-001", "filename": ..., "content": ...}]
    Les IDs sont attribués dans l'ordre de la liste (traçabilité).
    """
    documents: list[dict] = []
    for index, (filename, content) in enumerate(files, start=1):
        documents.append(
            {
                "id": f"DOC-{index:03d}",
                "filename": filename,
                "content": parse(filename, content, max_mb),
            }
        )
    return documents