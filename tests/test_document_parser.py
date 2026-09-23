"""Tests du parsing documentaire (TXT/MD/JSON/CSV/DOCX + refus des autres)."""

from io import BytesIO

import pytest

from app.services.document_parser import (
    DocumentParseError,
    load_documents,
    parse,
)


def test_parse_txt() -> None:
    assert parse("sys.txt", "Système interne".encode()) == "Système interne"


def test_parse_markdown() -> None:
    assert parse("archi.md", b"# Architecture\n- serveur linux") == (
        "# Architecture\n- serveur linux"
    )


def test_parse_json_as_text() -> None:
    payload = b'{"system": "web-app"}'
    assert parse("config.json", payload) == '{"system": "web-app"}'


def test_parse_csv_as_text() -> None:
    assert parse("users.csv", b"name,role") == "name,role"


def test_parse_pdf_invalid_content_fails_cleanly() -> None:
    with pytest.raises(DocumentParseError, match="PDF"):
        parse("doc.pdf", b"not a real pdf file")


def test_parse_docx_roundtrip() -> None:
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Application web interne")
    document.add_paragraph("Serveur Linux")
    document.save(buffer)

    text = parse("rapport.docx", buffer.getvalue())
    assert "Application web interne" in text
    assert "Serveur Linux" in text


def test_parse_rejects_unknown_extension() -> None:
    with pytest.raises(DocumentParseError, match="non autorisée"):
        parse("virus.exe", b"x")


def test_parse_rejects_oversized() -> None:
    with pytest.raises(DocumentParseError, match="volumineux"):
        parse("big.txt", b"x" * (11 * 1024 * 1024), max_mb=10)


def test_parse_rejects_invalid_utf8() -> None:
    with pytest.raises(DocumentParseError, match="encodage"):
        parse("fichier.txt", b"\xff\xfe invalid")


def test_load_documents_assigns_ids() -> None:
    documents = load_documents([("a.md", "système A".encode()), ("b.txt", "système B".encode())])
    assert [d["id"] for d in documents] == ["DOC-001", "DOC-002"]
    assert documents[0]["content"] == "système A"
    assert documents[0]["filename"] == "a.md"


def test_load_documents_rejects_bad_file() -> None:
    with pytest.raises(DocumentParseError):
        load_documents([("a.md", "système A".encode()), ("b.sh", b"#!/bin/bash")])